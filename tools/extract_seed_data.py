#!/usr/bin/env python3
"""
One-off extraction: pulls master data + history out of the two legacy Google
Sheets exports (PMS + Paint Consumption Tool) into clean CSVs consumed by
database/seeders/LegacyImportSeeder.php.

Usage:
  python3 tools/extract_seed_data.py <PMS.xlsx> <ConsumptionTool.xlsx> [outdir]
"""
import csv, re, sys
from collections import defaultdict
from datetime import datetime, timedelta
from openpyxl import load_workbook

pms_path, pct_path = sys.argv[1], sys.argv[2]
out = sys.argv[3] if len(sys.argv) > 3 else "database/seeders/data"

def pad(r, n=24):
    r = list(r); return r + [None] * (n - len(r))

def s(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def num(v):
    if v in (None, ""): return ""
    try: return f"{float(v):g}"
    except (TypeError, ValueError): return ""

def dt(v):
    if isinstance(v, datetime): return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, (int, float)):  # sheets serial
        return (datetime(1899, 12, 30) + timedelta(days=float(v))).strftime("%Y-%m-%d %H:%M:%S")
    return ""

def write(name, header, rows):
    with open(f"{out}/{name}", "w", newline="") as f:
        w = csv.writer(f); w.writerow(header); w.writerows(rows)
    print(f"{name}: {len(rows)} rows")

pms = load_workbook(pms_path, read_only=True, data_only=True)
pct = load_workbook(pct_path, read_only=True, data_only=True)

# ---------- bom_categories (Category_Map) ----------
cats = []
for r in map(pad, pms["Category_Map"].iter_rows(min_row=2, values_only=True)):
    if not r[0]: continue
    cats.append([s(r[0]), s(r[1]), 1 if s(r[2]).upper() == "YES" else 0, s(r[3]), s(r[4])])
write("bom_categories.csv", ["name", "issue_pool", "clubbed", "uom", "notes"], cats)

# ---------- items ----------
# name fallback + brand from the Inventory List sheet
inv_names = {}
for r in map(pad, pms["Paint Department Inventory List"].iter_rows(min_row=3, values_only=True)):
    if r[1] and r[2]: inv_names[s(r[1])] = (s(r[2]), s(r[4]), s(r[5]))
# density + odoo id from Consumption Tool Inventory (keyed by iFlow code)
density = {}
for r in map(pad, pct["Inventory"].iter_rows(min_row=6, values_only=True)):
    if r[0] and r[2]:
        density[s(r[2])] = (num(r[7]), s(r[1]), s(r[3]))

items, opening = [], {}
for r in map(pad, pms["Item_Master"].iter_rows(min_row=2, values_only=True)):
    code = s(r[0])
    if not code or code.lower() in ("brand",): continue
    name = s(r[1]) or inv_names.get(code, ("",))[0]
    dens, odoo, legacy_cat = density.get(code, ("", "", ""))
    items.append([code, odoo, name, s(r[3]), s(r[2]), s(r[4]), s(r[5]), s(r[7]), s(r[8]) or "Gram",
                  dens, num(r[16]), num(r[14]), legacy_cat, s(r[18])])
    opening[code] = (num(r[9]), num(r[13]))
write("items.csv", ["code", "odoo_id", "name", "brand", "manufacturer_code", "conventional_name",
                    "bom_category", "issue_pool", "uom", "density_kg_per_l", "rate_per_uom",
                    "min_level", "legacy_category", "pantone_ref"], items)
item_codes = {i[0] for i in items}

# ---------- orders + bom_lines ----------
orders = {}
for r in map(pad, pms["BOM_Allocation"].iter_rows(min_row=2, values_only=True)):
    if not s(r[0]): continue
    code = s(r[0]); bom = s(r[1])
    m = re.search(r"Total:\s*([\d,\.]+)", bom)
    total = m.group(1).replace(",", "") if m else ""
    orders[code] = [code, num(r[3]), num(r[4]), total, bom]

# Legacy sheet has the same order pasted more than once (last paste wins);
# BOM_Parsed therefore carries duplicate (order, category) lines — keep the last.
bom_by_key = {}
for r in map(pad, pms["BOM_Parsed"].iter_rows(min_row=2, values_only=True)):
    if not s(r[0]): continue
    bom_by_key[(s(r[0]), s(r[1]))] = [s(r[0]), s(r[1]), s(r[2]), num(r[3]) or "0", s(r[4]) or "Gram"]
    orders.setdefault(s(r[0]), [s(r[0]), "", "", "", ""])
bom_lines = list(bom_by_key.values())

# ---------- transactions (Stock_Ledger) ----------
txns = []
for r in map(pad, pms["Stock_Ledger"].iter_rows(min_row=2, values_only=True)):
    if not r[0]: continue
    typ = s(r[2]).upper()
    qin, qout = float(r[8] or 0), float(r[9] or 0)
    qty = qin - qout
    order_code = s(r[3])
    ref = ""
    if typ != "ISSUE" or order_code not in orders:
        # receipts/adjustments carry free-text refs, not order codes
        if order_code not in orders: ref, order_code = order_code, ""
    issue_type = {"BOM": "bom", "REWORK": "rework", "": ""}.get(s(r[14]).upper(), s(r[14]).lower())
    txns.append([s(r[0]), dt(r[1]), typ.lower(), issue_type, order_code, s(r[5]), s(r[4]),
                 f"{qty:g}", s(r[10]) or "Gram", num(r[17]), s(r[11]), s(r[12]), s(r[13]),
                 s(r[15]), s(r[16]), ref])
    if order_code and order_code not in orders:
        orders[order_code] = [order_code, "", "", "", ""]

# opening balances become transactions so stock is purely ledger-derived
first_ts = min(t[1] for t in txns if t[1])
open_ts = (datetime.strptime(first_ts, "%Y-%m-%d %H:%M:%S") - timedelta(days=1)).strftime("%Y-%m-%d 00:00:00")
for code, (open_qty, _) in opening.items():
    if open_qty and float(open_qty) != 0:
        txns.insert(0, [f"OPEN-{code}", open_ts, "opening", "", "", code, "", open_qty, "Gram", "", "system", "", "Opening stock from legacy Item_Master", "", "", ""])

write("transactions.csv", ["txn_ref", "occurred_at", "type", "issue_type", "order_code", "item_code",
                           "issue_pool", "qty", "uom", "rate", "entered_by", "authorized_by", "remarks",
                           "colour_ref", "batch_ref", "external_ref"], txns)

# ---------- colour batches ----------
batches, comps = [], []
for r in map(pad, pms["Colour_Log"].iter_rows(min_row=2, values_only=True)):
    if not r[0]: continue
    ref = s(r[0]); order_code = s(r[2])
    batches.append([ref, dt(r[1]), order_code, s(r[3]), s(r[5]), num(r[6]), num(r[7]), num(r[8]),
                    num(r[9]), s(r[12]), s(r[13]), s(r[14])])
    orders.setdefault(order_code, [order_code, "", "", "", ""])
    total = float(r[9] or 0) or 1
    for part in s(r[11]).split("|"):
        if ":" not in part: continue
        code, g = part.rsplit(":", 1)
        comps.append([ref, code.strip(), num(g), f"{float(g)/total*100:.2f}"])
write("colour_batches.csv", ["ref", "mixed_at", "order_code", "colour_ref", "hex", "lab_l", "lab_a",
                             "lab_b", "batch_grams", "matcher", "notes", "brand_mix"], batches)
write("colour_batch_components.csv", ["batch_ref", "item_code", "grams", "pct"], comps)
write("orders.csv", ["code", "total_area", "surface_area", "bom_total_cost", "bom_source"], list(orders.values()))
write("bom_lines.csv", ["order_code", "bom_category", "issue_pool", "allocated_qty", "uom"], bom_lines)

# ---------- validation: ledger-derived stock vs legacy Stock In Hand ----------
stock = defaultdict(float)
for t in txns:
    if t[5]: stock[t[5]] += float(t[7])
unknown = {t[5] for t in txns if t[5] and t[5] not in item_codes}
print("\nitems referenced in ledger but missing from Item_Master:", sorted(unknown) or "none")
mism = [(c, stock[c], float(sih or 0)) for c, (_, sih) in opening.items() if abs(stock[c] - float(sih or 0)) > 0.5]
print(f"stock mismatches vs legacy Stock In Hand: {len(mism)} of {len(opening)}")
for c, a, b in mism[:15]: print(f"  {c}: ledger={a:g} legacy={b:g}")
